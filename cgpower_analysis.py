"""
cgpower_analysis.py
Independent check of the CG Power research model + charts for the report.

Usage:
    python cgpower_analysis.py model/CGPower_Research_Model.xlsx

After you edit the workbook, open and save it in Excel first so the cached
formula values are up to date; this script reads those values.

Requires: openpyxl, numpy, matplotlib   (pip install openpyxl numpy matplotlib)
"""
import os
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

NAVY, LIGHT, ORANGE, GREY = "#1F3864", "#8EA9DB", "#C55A11", "#7F7F7F"
TOL = 0.5  # ₹ per share, or bps, tolerance for "matches"


# ---------------------------------------------------------------- helpers
def row_of(ws, label):
    """Row number whose column-A text equals label (so the script survives inserted rows)."""
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == label:
            return r
    raise KeyError(f"'{label}' not found in column A of sheet '{ws.title}'")


def header_cols(ws, header_row, labels):
    found = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if v in labels:
            found[v] = c
    return [found[x] for x in labels]


def row_values(ws, label, cols):
    r = row_of(ws, label)
    return np.array([ws.cell(r, c).value or 0.0 for c in cols], dtype=float)


def val(ws, label, col=2):
    return ws.cell(row_of(ws, label), col).value


def check(name, excel, python, tol=TOL):
    diff = float(np.max(np.abs(np.asarray(excel, dtype=float) - np.asarray(python, dtype=float))))
    status = "PASS" if diff <= tol else "FAIL"
    print(f"  [{status}] {name}: max difference {diff:,.4f}")
    return status == "PASS"


def grid(ws, anchor_label, nrows, ncols):
    r0 = row_of(ws, anchor_label)
    cols_hdr = [ws.cell(r0, 2 + j).value for j in range(ncols)]
    rows_hdr = [ws.cell(r0 + 1 + i, 1).value for i in range(nrows)]
    body = [[ws.cell(r0 + 1 + i, 2 + j).value for j in range(ncols)] for i in range(nrows)]
    return np.array(rows_hdr, float), np.array(cols_hdr, float), np.array(body, float)


# ---------------------------------------------------------------- load
path = sys.argv[1] if len(sys.argv) > 1 else os.path.join("model", "CGPower_Research_Model.xlsx")
wb = load_workbook(path, data_only=True)
H, I, FC, D, C, P = (wb[s] for s in ["Historicals", "Inputs", "Forecast", "DCF", "Commodity", "Comps"])

if D.cell(row_of(D, "DCF value per share (₹)"), 2).value is None:
    sys.exit("No cached values found. Open the workbook in Excel, save it, then rerun.")

price = val(I, "Share price, CMP (₹)")
shares = val(I, "Shares outstanding (Cr)")
net_cash = val(I, "Net cash (₹ Cr)")
wacc = val(I, "WACC")
g = val(I, "Terminal growth")

hist_years = ["FY22", "FY23", "FY24", "FY25", "FY26"]
fc_years = ["FY26A", "FY27E", "FY28E", "FY29E", "FY30E", "FY31E"]
hcols = header_cols(H, 4, hist_years)
fcols = header_cols(FC, 4, fc_years)

print(f"\nCG Power model check — {os.path.basename(path)}")
print(f"  Price ₹{price:,.0f} | WACC {wacc:.2%} | terminal growth {g:.2%} | net cash ₹{net_cash:,.0f} Cr\n")
ok = True

# ---------------------------------------------------------------- 1. forecast arithmetic
seg = sum(row_values(FC, lab, fcols) for lab in
          ["Power Systems revenue", "Industrial Systems revenue", "Semiconductor revenue", "Other / eliminations"])
ok &= check("Total revenue = sum of segments", row_values(FC, "Total revenue", fcols), seg)

# ---------------------------------------------------------------- 2. DCF rebuilt in Python
fcff = row_values(FC, "FCFF", fcols[1:])
rev = row_values(FC, "Total revenue", fcols)
g31 = rev[-1] / rev[-2] - 1
margin31 = fcff[-1] / rev[-1]
revs, flows = [rev[-1]], list(fcff)
for k in range(1, 6):
    gk = g31 - (g31 - g) * k / 5            # growth fades linearly to terminal growth
    revs.append(revs[-1] * (1 + gk))
    flows.append(revs[-1] * margin31)
flows = np.array(flows)
t = np.arange(1, 11)
pv = flows / (1 + wacc) ** t
tv = flows[-1] * (1 + g) / (wacc - g)
ev = pv.sum() + tv / (1 + wacc) ** 10
py_value = (ev + net_cash) / shares
xl_value = val(D, "DCF value per share (₹)")
ok &= check("DCF value per share", xl_value, py_value)

# ---------------------------------------------------------------- 3. WACC x g table
w_rows, g_cols, sens = grid(D, "WACC ↓ / terminal growth →", 5, 5)
py_sens = np.array([[((flows / (1 + w) ** t).sum() + flows[-1] * (1 + gg) / (w - gg) / (1 + w) ** 10 + net_cash) / shares
                     for gg in g_cols] for w in w_rows])
ok &= check("WACC × growth sensitivity table", sens, py_sens)

# ---------------------------------------------------------------- 4. reverse DCF, year by year (not closed form)
cagr_rows, m_cols, rev_dcf = grid(D, "10-yr revenue CAGR ↓ / FCFF margin →", 7, 4)
r0 = rev[0]


def steady_value(cagr, m):
    f = r0 * (1 + cagr) ** t * m
    return ((f / (1 + wacc) ** t).sum() + f[-1] * (1 + g) / (wacc - g) / (1 + wacc) ** 10 + net_cash) / shares


py_rev = np.array([[steady_value(c, m) for m in m_cols] for c in cagr_rows])
ok &= check("Reverse DCF table (closed-form Excel vs year-by-year Python)", rev_dcf, py_rev)

# ---------------------------------------------------------------- 5. commodity grid
rm, cu, st, pt = (val(C, x) for x in ["Raw material cost % of revenue", "Copper share of raw material cost",
                                      "Electrical steel share of raw material cost", "Price pass-through within the year"])
cu_rows, st_cols, bps = grid(C, "Copper change ↓ / Steel change →", 5, 5)
py_bps = np.array([[-rm * (cu * a + st * b) * (1 - pt) * 10000 for b in st_cols] for a in cu_rows])
ok &= check("Copper/steel margin grid (bps)", bps, py_bps)

print(f"\n  {'ALL CHECKS PASSED' if ok else 'SOME CHECKS FAILED — find the broken formula before presenting'}\n")

# ---------------------------------------------------------------- what the market is pricing in
print("What today's price implies (reverse DCF): FCFF margin needed for each 10-year revenue CAGR")
for c, vals in zip(cagr_rows, rev_dcf):
    slope = (vals[-1] - vals[0]) / (m_cols[-1] - m_cols[0])  # value per share is linear in the margin
    need = m_cols[0] + (price - vals[0]) / slope
    print(f"  {c:.0%} CAGR -> {need:.1%} FCFF margin")
print(f"  Base-case DCF ₹{xl_value:,.0f} vs price ₹{price:,.0f} ({xl_value / price - 1:+.0%})\n")

# ---------------------------------------------------------------- charts
os.makedirs("charts", exist_ok=True)
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.spines.top": False,
                     "axes.spines.right": False})

# (a) revenue and EBITDA margin
h_rev = row_values(H, "Revenue", hcols)
h_m = row_values(H, "EBITDA margin", hcols)
f_rev = rev[1:]
f_m = row_values(FC, "EBITDA margin", fcols[1:])
labels = hist_years + fc_years[1:]
fig, ax = plt.subplots(figsize=(9, 4.5))
ax.bar(labels, np.r_[h_rev, f_rev], color=[NAVY] * 5 + [LIGHT] * 5)
ax.set_ylabel("Revenue (₹ Cr)")
ax2 = ax.twinx()
ax2.plot(labels, np.r_[h_m, f_m] * 100, color=ORANGE, marker="o")
ax2.set_ylabel("EBITDA margin (%)", color=ORANGE)
ax2.spines["top"].set_visible(False)
ax.set_title("Revenue and EBITDA margin: FY22–FY26 actual, FY27E–FY31E forecast")
fig.tight_layout()
fig.savefig("charts/1_revenue_margin.png", dpi=200)
plt.close(fig)

# (b) segment mix FY26 vs FY31E
segs = ["Power Systems revenue", "Industrial Systems revenue", "Semiconductor revenue"]
names = ["Power Systems", "Industrial Systems\n(motors, drives, rail)", "Semiconductor"]
v26 = [row_values(FC, s, fcols[:1])[0] for s in segs]
v31 = [row_values(FC, s, fcols[-1:])[0] for s in segs]
x = np.arange(3)
fig, ax = plt.subplots(figsize=(8, 4.2))
ax.bar(x - 0.2, v26, 0.4, label="FY26A", color=NAVY)
ax.bar(x + 0.2, v31, 0.4, label="FY31E", color=LIGHT)
ax.set_xticks(x, names)
ax.set_ylabel("₹ Cr")
ax.legend(frameon=False)
ax.set_title("Where growth comes from: segment revenue FY26A vs FY31E")
fig.tight_layout()
fig.savefig("charts/2_segment_mix.png", dpi=200)
plt.close(fig)


def heatmap(data, rows, cols, rfmt, cfmt, title, fname, cbar, fmt="₹{:,.0f}", highlight=None, cmap="RdYlGn"):
    fig, ax = plt.subplots(figsize=(7.5, 4.6))
    im = ax.imshow(data, cmap=cmap, aspect="auto")
    ax.set_xticks(range(len(cols)), [cfmt(c) for c in cols])
    ax.set_yticks(range(len(rows)), [rfmt(r) for r in rows])
    for i in range(data.shape[0]):
        for j in range(data.shape[1]):
            bold = highlight is not None and data[i, j] >= highlight
            ax.text(j, i, fmt.format(data[i, j]), ha="center", va="center", fontsize=9,
                    fontweight="bold" if bold else "normal")
    fig.colorbar(im, ax=ax, label=cbar)
    ax.set_title(title)
    fig.tight_layout()
    fig.savefig(fname, dpi=200)
    plt.close(fig)


pct = lambda v: f"{v:.1%}"
heatmap(sens, w_rows, g_cols, pct, pct, "DCF value per share: WACC (rows) vs terminal growth (columns)",
        "charts/3_dcf_sensitivity.png", "₹ per share")
heatmap(rev_dcf, cagr_rows, m_cols, lambda v: f"{v:.0%} CAGR", lambda v: f"{v:.0%} margin",
        f"What ₹{price:,.0f} implies: value/share by 10-yr revenue CAGR and FCFF margin\n(bold = at or above today's price)",
        "charts/4_reverse_dcf.png", "₹ per share", highlight=price)
heatmap(bps, cu_rows, st_cols, lambda v: f"Copper {v:+.0%}", lambda v: f"Steel {v:+.0%}",
        "EBITDA margin impact (bps) from copper and electrical-steel moves", "charts/5_commodity.png",
        "basis points", fmt="{:+,.0f}")

# (c) valuation summary ("football field")
bars = [("DCF (WACC ± 1pp, g ± 1pp)", sens.min(), sens.max())]
comps_val = P.cell(row_of(P, "Comps value per share (₹)"), 2).value
if isinstance(comps_val, (int, float)):
    lo = P.cell(row_of(P, "At peer median P/E (₹/share)"), 2).value
    hi = P.cell(row_of(P, "At peer median EV/EBITDA (₹/share)"), 2).value
    vals_ = [v for v in (lo, hi) if isinstance(v, (int, float))]
    bars.append(("Peer multiples (P/E, EV/EBITDA)", min(vals_), max(vals_)))
bars.append((f"Scenarios ({cagr_rows.min():.0%}–{cagr_rows.max():.0%} CAGR, {m_cols.min():.0%}–{m_cols.max():.0%} margin)",
             rev_dcf.min(), rev_dcf.max()))
fig, ax = plt.subplots(figsize=(8.5, 3.2))
for k, (name, lo, hi) in enumerate(bars):
    ax.barh(k, hi - lo, left=lo, color=LIGHT, edgecolor=NAVY)
    ax.text(hi, k, f"  ₹{lo:,.0f}–₹{hi:,.0f}", va="center", fontsize=9)
ax.axvline(price, color=ORANGE, linestyle="--")
ax.text(price, -0.55, f" Price ₹{price:,.0f}", color=ORANGE, fontsize=9, va="bottom")
ax.set_yticks(range(len(bars)), [b[0] for b in bars])
ax.invert_yaxis()
ax.set_xlabel("₹ per share")
ax.set_xlim(left=0)
ax.set_title("Valuation summary", pad=14)
fig.tight_layout()
fig.savefig("charts/6_valuation_summary.png", dpi=200)
plt.close(fig)

# (d) peer multiples, once the Comps sheet is filled
r0p = row_of(P, "CG Power (subject)")
peer_rows = [(P.cell(r, 1).value, P.cell(r, 9).value, P.cell(r, 10).value) for r in range(r0p, r0p + 6)]
peer_rows = [p for p in peer_rows if isinstance(p[1], (int, float)) and isinstance(p[2], (int, float))]
if len(peer_rows) > 1:
    fig, axes = plt.subplots(1, 2, figsize=(10, 4))
    for ax, idx, name in [(axes[0], 1, "EV/EBITDA (FY26)"), (axes[1], 2, "P/E (FY26)")]:
        ax.barh([p[0] for p in peer_rows], [p[idx] for p in peer_rows],
                color=[ORANGE if "CG Power" in p[0] else LIGHT for p in peer_rows])
        ax.set_title(name)
        ax.invert_yaxis()
    fig.tight_layout()
    fig.savefig("charts/7_peer_multiples.png", dpi=200)
    plt.close(fig)
    print("Charts saved to ./charts (including peer multiples).")
else:
    print("Charts saved to ./charts (fill the Comps sheet to add the peer-multiples chart).")
